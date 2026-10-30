from rest_framework import viewsets, status
from rest_framework.response import Response
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from .models import Batch, Student, Attendance
from .serializers import BatchSerializer, StudentSerializer, AttendanceSerializer, BulkAttendanceSerializer
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.db.models import Count, Q
from django.utils import timezone
from datetime import datetime

class BatchViewSet(viewsets.ModelViewSet):
    queryset = Batch.objects.all().order_by('-created_at')
    serializer_class = BatchSerializer
    permission_classes = [IsAuthenticated]

class StudentViewSet(viewsets.ModelViewSet):
    queryset = Student.objects.all().order_by('student_id')
    serializer_class = StudentSerializer
    permission_classes = [IsAuthenticated]

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_attendance(request):
    batch_id = request.query_params.get('batch')
    date = request.query_params.get('date')
    
    if not batch_id or not date:
        return Response({'error': 'batch and date parameters are required'}, status=status.HTTP_400_BAD_REQUEST)
        
    attendances = Attendance.objects.filter(batch_id=batch_id, date=date)
    serializer = AttendanceSerializer(attendances, many=True)
    return Response(serializer.data)

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def save_bulk_attendance(request):
    serializer = BulkAttendanceSerializer(data=request.data)
    if serializer.is_valid():
        batch = serializer.validated_data['batch']
        date = serializer.validated_data['date']
        records = serializer.validated_data['records']
        
        Attendance.objects.filter(batch=batch, date=date).delete()
        
        attendance_objects = []
        for record in records:
            student_id = record.get('student')
            attn_status = record.get('status', 'PRESENT')
            
            attendance_objects.append(Attendance(
                student_id=student_id,
                batch=batch,
                date=date,
                status=attn_status
            ))
            
        Attendance.objects.bulk_create(attendance_objects)
        return Response({'message': 'Attendance saved successfully'}, status=status.HTTP_200_OK)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_stats(request):
    today = timezone.now().date()
    batches = Batch.objects.all()
    stats = []
    
    for batch in batches:
        total_students = Student.objects.filter(batch=batch).count()
        todays_attendance = Attendance.objects.filter(batch=batch, date=today)
        
        present = todays_attendance.filter(status='PRESENT').count()
        absent = todays_attendance.filter(status='ABSENT').count()
        
        stats.append({
            'batch_id': batch.id,
            'batch_name': batch.name,
            'total_students': total_students,
            'present': present,
            'absent': absent,
            'is_marked': todays_attendance.exists()
        })
        
    return Response({
        'date': today,
        'stats': stats
    })

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def batch_report(request):
    batch_id = request.query_params.get('batch')
    
    if not batch_id:
        return Response({'error': 'batch parameter is required'}, status=status.HTTP_400_BAD_REQUEST)
        
    try:
        batch = Batch.objects.get(id=batch_id)
    except Batch.DoesNotExist:
        return Response({'error': 'Batch not found'}, status=status.HTTP_404_NOT_FOUND)
        
    students = Student.objects.filter(batch=batch).order_by('student_id')
    attendances = Attendance.objects.filter(batch=batch).order_by('date')
    
    dates = list(attendances.values_list('date', flat=True).distinct().order_by('date'))
    date_strings = [d.strftime('%d-%b') for d in dates]
    
    report_data = []
    
    for student in students:
        student_att = attendances.filter(student=student)
        present_count = 0
        absent_count = 0
        attendance_record = {}
        
        for d, d_str in zip(dates, date_strings):
            att = student_att.filter(date=d).first()
            if att:
                status_code = 'P' if att.status == 'PRESENT' else 'A'
                if att.status == 'PRESENT': present_count += 1
                else: absent_count += 1
                attendance_record[d_str] = status_code
            else:
                attendance_record[d_str] = '-'
                
        report_data.append({
            'student_id': student.student_id,
            'name': student.name,
            'attendance': attendance_record,
            'present': present_count,
            'absent': absent_count
        })
        
    return Response({
        'dates': date_strings,
        'students': report_data
    })

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_excel(request):
    batch_id = request.query_params.get('batch')
    
    if not batch_id:
        return Response({'error': 'batch parameter is required'}, status=status.HTTP_400_BAD_REQUEST)
        
    try:
        batch = Batch.objects.get(id=batch_id)
    except Batch.DoesNotExist:
        return Response({'error': 'Batch not found'}, status=status.HTTP_404_NOT_FOUND)
        
    students = Student.objects.filter(batch=batch).order_by('student_id')
    attendances = Attendance.objects.filter(batch=batch).order_by('date')
    
    dates = list(attendances.values_list('date', flat=True).distinct().order_by('date'))
    
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = f"{batch.name} Attendance"
    
    headers = ['Student ID', 'Name'] + [d.strftime('%d-%b') for d in dates] + ['Present', 'Absent']
    sheet.append(headers)
    
    # Define Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5") # Match primary color
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    present_font = Font(color="10B981", bold=True)
    absent_font = Font(color="EF4444", bold=True)
    
    # Apply Header Styles
    for col_num, cell in enumerate(sheet[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        
        # Set column widths
        col_letter = get_column_letter(col_num)
        if col_num == 1:
            sheet.column_dimensions[col_letter].width = 15
        elif col_num == 2:
            sheet.column_dimensions[col_letter].width = 25
        else:
            sheet.column_dimensions[col_letter].width = 8

    sheet.freeze_panes = 'C2' # Freeze ID, Name and Header row

    row_idx = 2
    for student in students:
        row = [student.student_id, student.name]
        student_att = attendances.filter(student=student)
        
        present_count = 0
        absent_count = 0
        
        for d in dates:
            att = student_att.filter(date=d).first()
            if att:
                status_code = 'P' if att.status == 'PRESENT' else 'A'
                if att.status == 'PRESENT': present_count += 1
                else: absent_count += 1
                row.append(status_code)
            else:
                row.append('-')
                
        total_classes = present_count + absent_count
        
        row.extend([present_count, absent_count])
        sheet.append(row)
        
        # Apply Row Styles
        for col_num, cell in enumerate(sheet[row_idx], 1):
            cell.border = thin_border
            if col_num <= 2:
                cell.alignment = left_align
            else:
                cell.alignment = center_align
                
                if cell.value == 'P':
                    cell.font = present_font
                elif cell.value == 'A':
                    cell.font = absent_font
                    
        row_idx += 1
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="attendance_{batch.name}.xlsx"'
    workbook.save(response)
    
    return response

